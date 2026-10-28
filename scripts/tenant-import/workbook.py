"""Private XLSX ingestion: data only; no evaluation/network/macro execution."""
import hashlib,json,zipfile
from collections import Counter,defaultdict
from pathlib import Path
import openpyxl
PARSER_VERSION='teams-one-time-1'
SOURCE_SHA='8a6dcf40cb6f8d7d66322c0048e52f930615d266674cfbf1ecacd5e9cc024311'
HEADERS={
'Import Review':'record_id decision property_map_key property_label_proposed existing_property_id unit_label_proposed existing_unit_id tenant_name_proposed existing_tenant_id source_customer_raw source_sheet source_row repeat_group repeat_group_size review_reasons reviewer review_note'.split(),
'Contacts':'record_id contact_type source_role source_header_raw contact_value_raw normalized_candidate country_confirmed validation_state source_sheet source_cell decision'.split(),
'Property Map':'property_map_key source_sheet source_property_alias property_label_proposed existing_property_id confirmed_address decision review_note'.split(),
'Source Cells':'source_sheet source_cell source_type source_value_raw'.split()}
IMMUTABLE={'Import Review':'record_id source_customer_raw source_sheet source_row property_map_key'.split(),'Contacts':'record_id source_role source_header_raw contact_value_raw source_sheet source_cell'.split(),'Property Map':'property_map_key source_sheet source_property_alias'.split()}
def canonical(value):
 if isinstance(value,dict):return {k:canonical(v) for k,v in value.items()}
 if isinstance(value,list):return [canonical(v) for v in value]
 if isinstance(value,float) and value.is_integer():return int(value)
 return value
def digest(value):return hashlib.sha256(json.dumps(canonical(value),sort_keys=True,separators=(',',':'),ensure_ascii=False).encode()).hexdigest()
def filehash(path):return hashlib.sha256(Path(path).read_bytes()).hexdigest()
def require(ok,message):
 if not ok:raise ValueError(message)
def key(sheet,row):
 if sheet=='Import Review':return row['record_id']
 if sheet=='Contacts':return (row['record_id'],row['source_sheet'],row['source_cell'])
 return row['property_map_key']
def load(path):
 require(Path(path).stat().st_size<20_000_000,'Workbook too large')
 with zipfile.ZipFile(path) as z:
  require(sum(i.file_size for i in z.infolist())<100_000_000,'Expanded workbook too large')
  require(not any('vbaproject' in n.lower() or 'externallinks/' in n.lower() for n in z.namelist()),'Macros/external links are not allowed')
 w=openpyxl.load_workbook(path,read_only=True,data_only=False,keep_links=False)
 require(set(w.sheetnames)==set(HEADERS)|{'Start Here'},'Unexpected/missing worksheets')
 out={}
 for name,headers in HEADERS.items():
  rows=iter(w[name].iter_rows());head=[c.value for c in next(rows)]
  require(len(head)==len(headers) and len(set(head))==len(head) and set(head)==set(headers),f'Unexpected/duplicate headers in {name}')
  result=[]
  for cells in rows:
   if not any(c.value is not None for c in cells):continue
   require(not any(c.data_type=='f' for c in cells),f'Formula cell in {name}')
   require(len(cells)<=len(head) or not any(c.value is not None for c in cells[len(head):]),f'Unexpected columns in {name}')
   result.append({h:(cells[i].value if i<len(cells) and cells[i].value is not None else '') for i,h in enumerate(head)})
  out[name]=result
 for name in IMMUTABLE:
  keys=[key(name,r) for r in out[name]];require(len(keys)==len(set(keys)),f'Duplicate keys in {name}')
 contact_keys=[(r['source_sheet'],r['source_cell']) for r in out['Contacts']]
 require(len(contact_keys)==len(set(contact_keys)),'Each original contact cell must be accounted for once')
 source_keys=[(r['source_sheet'],r['source_cell']) for r in out['Source Cells']]
 require(len(source_keys)==len(set(source_keys)),'Duplicate source cells')
 start={r[0].value:r[1].value for r in w['Start Here'].iter_rows() if len(r)>1}
 require(start.get('Source SHA-256')==SOURCE_SHA,'Wrong source hash')
 out['source_sha256']=SOURCE_SHA;out['workbook_sha256']=filehash(path);out['parser_version']=PARSER_VERSION
 return out

def validate_against(package,baseline):
 for field in ['source_sha256','parser_version']:require(package[field]==baseline[field],f'Changed {field}')
 require(package['Source Cells']==baseline['Source Cells'],'Immutable raw source archive changed')
 for name,fields in IMMUTABLE.items():
  original={key(name,r):r for r in baseline[name]};current={key(name,r):r for r in package[name]}
  require(set(original)==set(current),f'Unknown/missing keys in {name}')
  for k,r in current.items():require(all(r[f]==original[k][f] for f in fields),f'Immutable provenance changed in {name}')
 return package

def report(p):
 parents={r['record_id']:r for r in p['Import Review']};cells={(r['source_sheet'],r['source_cell']) for r in p['Source Cells']}
 require(len(parents)==407 and len(p['Contacts'])==959 and len(cells)==1510 and len(p['Property Map'])==19,'Control totals mismatch')
 require(all(c['record_id'] in parents for c in p['Contacts']),'Unknown contact parent')
 require(all((c['source_sheet'],c['source_cell']) in cells for c in p['Contacts']),'Contact missing raw source cell')
 signatures=defaultdict(list)
 for r in p['Import Review']:
  contacts=[c['contact_value_raw'] for c in p['Contacts'] if c['record_id']==r['record_id']]
  sig=(' '.join(str(r['source_customer_raw']).split()).casefold(),tuple(' '.join(str(c).split()).casefold() for c in contacts))
  signatures[sig].append(r)
 repeats=[g for g in signatures.values() if len(g)>1]
 return {'source_sha256':p['source_sha256'],'workbook_sha256':p['workbook_sha256'],'parser_version':PARSER_VERSION,
 'candidates':len(parents),'mapped_contact_cells':len(p['Contacts']),'source_cells':len(cells),'property_aliases':len(p['Property Map']),
 'candidate_decisions':dict(Counter(r['decision'] for r in parents.values())), 'contact_decisions':dict(Counter(r['decision'] for r in p['Contacts'])),
 'property_decisions':dict(Counter(r['decision'] for r in p['Property Map'])),
 'recalculated_repeat_groups':len(repeats),'recalculated_excess_occurrences':sum(len(g)-1 for g in repeats),
 'imported':0,'remote_writes_enabled':False,'source_snapshot_independently_verified':False}
